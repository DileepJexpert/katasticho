"""Generate the React HSN lookup dataset from GSTN's official workbook.

Usage:
    python scripts/generate_react_hsn_directory.py HSN_SAC.xlsx \
        react_app/public/data/india-hsn-directory.json
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


VALID_CODE_LENGTHS = {2, 4, 6, 8}


def normalize_code(value: object) -> str:
    code = re.sub(r"\s+", "", str(value or ""))
    if len(code) == 5:
        code = code.zfill(6)
    elif len(code) == 7:
        code = code.zfill(8)
    return code


def normalize_description(value: object) -> str:
    description = str(value or "")
    description = description.replace("_x000D_", " ").replace("~", " - ")
    description = description.replace("\ufffd", "'")
    return re.sub(r"\s+", " ", description).strip()


def generate(source: Path, destination: Path) -> int:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["HSN_MSTR"]
    records: list[dict[str, str]] = []
    seen: set[str] = set()

    for code_value, description_value in sheet.iter_rows(
        min_row=2,
        max_col=2,
        values_only=True,
    ):
        code = normalize_code(code_value)
        if not code.isdigit() or len(code) not in VALID_CODE_LENGTHS or code in seen:
            continue

        seen.add(code)
        records.append(
            {
                "hsnCode": code,
                "description": normalize_description(description_value),
            }
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": "https://tutorial.gst.gov.in/downloads/HSN_SAC.xlsx",
        "retrievedOn": "2026-09-06",
        "records": records,
    }
    destination.write_text(
        json.dumps(payload, ensure_ascii=True, separators=(",", ":")),
        encoding="utf-8",
    )
    return len(records)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()

    count = generate(args.source, args.destination)
    print(f"Generated {count} HSN entries at {args.destination}")


if __name__ == "__main__":
    main()
